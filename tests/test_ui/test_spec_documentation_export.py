"""A profile exports as documentation of its model, in Markdown and Excel.

The explorer's export with one profile open was a comparison of the profile
with itself: entity names marked ``=`` and an empty "Changed Fields" section,
nothing about what a field means or holds. The documentation walks the
entities root first and lists every field with its type, requirement,
description, constraints and example.
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from metaseed.specs.loader import SpecLoader
from metaseed.specs.ordering import entity_order
from metaseed.ui.app import create_app
from metaseed.ui.services.spec_documentation import profile_markdown, profile_workbook
from metaseed.ui.state import AppState


@pytest.fixture
def spec():
    return SpecLoader().load_profile(version="1.0", profile="ena")


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app(AppState()))


def test_the_markdown_walks_the_entities_root_first(spec) -> None:
    text = profile_markdown(spec)

    headings = [line[4:] for line in text.splitlines() if line.startswith("### ")]
    assert headings[: len(spec.entities)] == entity_order(spec)
    assert headings[0] == spec.root_entity


def test_each_field_is_listed_with_its_description_constraints_and_example(
    spec,
) -> None:
    text = profile_markdown(spec)
    experiment = text.split("### Experiment\n", 1)[1].split("\n### ", 1)[0]

    row = next(
        line
        for line in experiment.splitlines()
        if line.startswith("| library_layout |")
    )
    assert "Library layout (SINGLE or PAIRED)." in row
    assert "SINGLE, PAIRED" in row
    assert "| yes |" in row
    field = next(f for f in spec.entities["Sample"].fields if f.name == "accession")
    assert field.constraints is not None and field.constraints.pattern
    assert field.constraints.pattern.replace("|", "\\|") in text


def test_the_workbook_describes_the_model_sheet_by_sheet(spec) -> None:
    book = load_workbook(io.BytesIO(profile_workbook(spec)))

    assert book.sheetnames == ["Entities", "Fields", "Rules"]
    entities = [
        row[0] for row in book["Entities"].iter_rows(min_row=2, values_only=True)
    ]
    assert entities == entity_order(spec)
    fields = list(book["Fields"].iter_rows(values_only=True))
    header = fields[0]
    assert header[:4] == ("Entity", "Field", "Type", "Required")
    layout = next(row for row in fields if row[:2] == ("Experiment", "library_layout"))
    assert "SINGLE, PAIRED" in layout[header.index("Constraints")]
    assert book["Fields"].freeze_panes == "A2"
    assert len(list(book["Rules"].iter_rows(min_row=2))) == len(spec.validation_rules)


def test_the_explorer_exports_one_profile_as_its_documentation(client) -> None:
    markdown = client.get("/explore/report/markdown/ena/1.0")
    assert markdown.status_code == 200
    assert "### Experiment" in markdown.text
    assert "Changed Fields" not in markdown.text

    workbook = client.get("/explore/report/xlsx/ena/1.0")
    assert workbook.status_code == 200
    assert "ena-1.0" in workbook.headers["content-disposition"]
    assert load_workbook(io.BytesIO(workbook.content)).sheetnames[0] == "Entities"


def test_two_profiles_still_export_their_comparison(client) -> None:
    response = client.get("/explore/report/markdown/ena/1.0,isa/1.0")
    assert "Profile Comparison Report" in response.text


def test_the_spec_builder_exports_its_draft_as_documentation(client) -> None:
    client.get("/spec-builder/clone/ena/1.0")

    markdown = client.get("/spec-builder/export/doc/markdown")
    assert markdown.status_code == 200
    assert "### Experiment" in markdown.text
    workbook = client.get("/spec-builder/export/doc/xlsx")
    assert load_workbook(io.BytesIO(workbook.content)).sheetnames == [
        "Entities",
        "Fields",
        "Rules",
    ]


def test_both_pages_offer_the_documentation() -> None:
    from metaseed.ui.app import TEMPLATES_DIR

    explore = (TEMPLATES_DIR / "explore" / "index.html").read_text()
    builder = (TEMPLATES_DIR / "spec_builder" / "base.html").read_text()
    assert "exportReport('xlsx')" in explore
    assert "/spec-builder/export/doc/markdown" in builder
    assert "/spec-builder/export/doc/xlsx" in builder


def test_the_markdown_downloads_as_a_file(client) -> None:
    """Served inline, the browser showed the Markdown as text in a tab."""
    response = client.get("/explore/report/markdown/ena/1.0")

    assert (
        response.headers["content-disposition"] == 'attachment; filename="ena-1.0.md"'
    )


def test_csv_of_one_profile_is_its_fields_table(client, spec) -> None:
    """CSV once returned the Markdown, under a CSV button."""
    import csv

    response = client.get("/explore/report/csv/ena/1.0")

    assert (
        response.headers["content-disposition"] == 'attachment; filename="ena-1.0.csv"'
    )
    assert response.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(response.text)))
    assert rows[0][:4] == ["Entity", "Field", "Type", "Required"]
    assert len(rows) - 1 == sum(len(e.fields) for e in spec.entities.values())


def test_one_profile_has_no_html_export() -> None:
    """HTML is a comparison report only; with one profile its button hides."""
    from metaseed.ui.app import TEMPLATES_DIR

    explore = (TEMPLATES_DIR / "explore" / "index.html").read_text()
    assert 'id="export-html"' in explore
    assert "getElementById('export-html').style.display = 'none'" in explore
