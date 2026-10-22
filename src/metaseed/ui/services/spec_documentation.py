"""A profile's model as documentation: Markdown and an Excel workbook.

Both walk the entities in containment order (root first) and describe every
field: type, requirement, description, constraints and the example value the
specification gives. The constraints text is the note a table heading shows on
hover (:func:`heading_note`), so the documentation and the editor describe a
field in the same words.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from metaseed.specs.ordering import child_entity_types, entity_order
from metaseed.specs.schema import (
    EntityDefSpec,
    FieldSpec,
    ProfileSpec,
    applies_to_entity,
)
from metaseed.ui.services.heading_note import heading_note

FIELD_COLUMNS = (
    "Entity",
    "Field",
    "Type",
    "Required",
    "Description",
    "Constraints",
    "Example",
    "Ontology term",
)
RULE_COLUMNS = ("Rule", "Type", "Applies to", "Field", "Description")
ENTITY_COLUMNS = (
    "Entity",
    "Description",
    "Ontology term",
    "Contained in",
    "Contains",
    "Fields",
)


def field_type(field: FieldSpec) -> str:
    """The field's type as written in the specification, with its item type."""
    kind = field.type.value if hasattr(field.type, "value") else str(field.type)
    return f"{kind} of {field.items}" if field.items else kind


def field_constraints(entity_name: str, field: FieldSpec, spec: ProfileSpec) -> str:
    """What a valid value is: the heading note without its description.

    Allowed values are listed in full; the heading note shortens a long list.
    """
    rules = [
        r for r in spec.validation_rules if applies_to_entity(r.applies_to, entity_name)
    ]
    note = heading_note(field.name, field, rules)
    description = (field.description or "").strip()
    if description and note.startswith(description):
        note = note[len(description) :]
    lines = [line for line in note.strip().splitlines() if line.strip()]
    # The first line is the "Required · Text" summary the Type and Required
    # columns already carry.
    lines = lines[1:]
    allowed = field.options or (field.constraints.enum if field.constraints else None)
    if allowed:
        lines = [
            f"Allowed values: {', '.join(map(str, allowed))}"
            if line.startswith("Allowed values:")
            else line
            for line in lines
        ]
    return "; ".join(lines)


def _example(entity: EntityDefSpec, field: FieldSpec) -> str:
    """The field's example value; nested entities are described by their own."""
    value: Any = field.example
    if value is None and isinstance(entity.example, dict):
        value = entity.example.get(field.name)
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(item) for item in value if not isinstance(item, dict))
    return "" if isinstance(value, dict) else str(value)


def _parents(spec: ProfileSpec) -> dict[str, list[str]]:
    parents: dict[str, list[str]] = {}
    for parent, children in child_entity_types(spec).items():
        for child in children:
            parents.setdefault(child, []).append(parent)
    return parents


def _cell(text: Any) -> str:
    """Text safe inside a Markdown table cell."""
    return str(text or "").replace("|", "\\|").replace("\n", " ").strip()


def profile_markdown(spec: ProfileSpec) -> str:
    """The profile's model as a Markdown document.

    Args:
        spec: The profile to document.

    Returns:
        Markdown: an overview, one section per entity in containment order
        with a table of its fields, and the validation rules.
    """
    order = entity_order(spec)
    children = child_entity_types(spec)
    parents = _parents(spec)
    title = spec.display_name or spec.name
    out = [f"# {title} {spec.version}", ""]
    if spec.description:
        out += [spec.description.strip(), ""]
    out += [
        f"- **Profile:** `{spec.name}` version `{spec.version}`",
        f"- **Root entity:** {spec.root_entity}",
        f"- **Entities:** {len(spec.entities)}, **fields:** "
        f"{sum(len(e.fields) for e in spec.entities.values())}, "
        f"**validation rules:** {len(spec.validation_rules)}",
        "",
        "## Entities",
        "",
    ]
    for name in order:
        entity = spec.entities[name]
        out += [f"### {name}", ""]
        if entity.description:
            out += [entity.description.strip(), ""]
        facts = []
        if entity.ontology_term:
            facts.append(f"- **Ontology term:** {entity.ontology_term}")
        if parents.get(name):
            facts.append(f"- **Contained in:** {', '.join(parents[name])}")
        if children.get(name):
            facts.append(f"- **Contains:** {', '.join(children[name])}")
        if facts:
            out += [*facts, ""]
        out += [
            "| Field | Type | Required | Description | Constraints | Example |",
            "| --- | --- | --- | --- | --- | --- |",
        ]
        for field in entity.fields:
            out.append(
                f"| {field.name} | {_cell(field_type(field))} | "
                f"{'yes' if field.required else 'no'} | {_cell(field.description)} | "
                f"{_cell(field_constraints(name, field, spec))} | "
                f"{_cell(_example(entity, field))} |"
            )
        out.append("")
    if spec.validation_rules:
        out += [
            "## Validation rules",
            "",
            "| Rule | Type | Applies to | Field | Description |",
            "| --- | --- | --- | --- | --- |",
        ]
        for rule in spec.validation_rules:
            out.append(
                f"| {_cell(rule.name)} | {_cell(rule.type)} | {_cell(_applies(rule))} | "
                f"{_cell(rule.field)} | {_cell(rule.description or rule.message)} |"
            )
        out.append("")
    return "\n".join(out)


def _applies(rule: Any) -> str:
    target = rule.applies_to
    return ", ".join(target) if isinstance(target, list) else str(target or "")


def _sheet(
    book: Workbook, title: str, header: tuple[str, ...], rows: list[list[Any]]
) -> None:
    sheet = book.create_sheet(title)
    sheet.append(list(header))
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    wrap = Alignment(wrap_text=True, vertical="top")
    for index, name in enumerate(header, start=1):
        long = name in ("Description", "Constraints", "Example")
        sheet.column_dimensions[get_column_letter(index)].width = 50 if long else 20
        if long:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.alignment = wrap


def _field_rows(spec: ProfileSpec) -> list[list[Any]]:
    """One row per field, entities in containment order, as ``FIELD_COLUMNS``."""
    return [
        [
            name,
            field.name,
            field_type(field),
            "yes" if field.required else "no",
            (field.description or "").strip(),
            field_constraints(name, field, spec),
            _example(spec.entities[name], field),
            field.ontology_term or "",
        ]
        for name in entity_order(spec)
        for field in spec.entities[name].fields
    ]


def profile_fields_csv(spec: ProfileSpec) -> str:
    """The profile's fields as CSV: the workbook's *Fields* sheet.

    Args:
        spec: The profile to document.

    Returns:
        CSV text with a header row of ``FIELD_COLUMNS``.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(FIELD_COLUMNS)
    writer.writerows(_field_rows(spec))
    return buffer.getvalue()


def profile_workbook(spec: ProfileSpec) -> bytes:
    """The profile's model as an Excel workbook.

    Args:
        spec: The profile to document.

    Returns:
        The ``.xlsx`` bytes: sheets *Entities*, *Fields* and *Rules*.
    """
    order = entity_order(spec)
    children = child_entity_types(spec)
    parents = _parents(spec)
    book = Workbook()
    book.remove(book.active)
    _sheet(
        book,
        "Entities",
        ENTITY_COLUMNS,
        [
            [
                name,
                (spec.entities[name].description or "").strip(),
                spec.entities[name].ontology_term or "",
                ", ".join(parents.get(name, [])),
                ", ".join(children.get(name, [])),
                len(spec.entities[name].fields),
            ]
            for name in order
        ],
    )
    _sheet(book, "Fields", FIELD_COLUMNS, _field_rows(spec))
    _sheet(
        book,
        "Rules",
        RULE_COLUMNS,
        [
            [
                rule.name,
                str(rule.type),
                _applies(rule),
                rule.field or "",
                (rule.description or rule.message or "").strip(),
            ]
            for rule in spec.validation_rules
        ],
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
