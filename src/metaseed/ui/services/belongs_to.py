"""The workbook's **Belongs to** column: where a row sits in the dataset's tree.

A row names the record that holds it as ``Type: identifier``, for example
``Study: STU-1``, and as ``Type.field: identifier`` when that type holds the
row's entity type in more than one field. One column replaces the earlier pair
of an identifier alone (``_parent``), which could not tell an investigation's
contact from a study's person, and a second dropdown naming the holding field.

The export writes the value for every stored row, the dropdown lists every
record that may hold a new one, and the import reads it back into
``_parent_unique_id`` and ``_parent_field``. See
docs/architecture/excel-round-trip.md.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Any, NamedTuple

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

#: The column's heading.
BELONGS_TO = "Belongs to"

#: The column workbooks exported before 0.67.0 carry: the parent's identifier.
LEGACY_PARENT = "_parent"

#: The hidden sheet listing, per entity type, every record that may hold a row.
PARENTS_SHEET = "metaseed parents"

#: Rows of each holder's sheet the dropdown reads, as the other dropdowns do.
PARENT_ROWS = 2000


class Holder(NamedTuple):
    """A field of one entity type that holds rows of another."""

    entity_type: str
    field: str


class Placement(NamedTuple):
    """Where a row goes: the holding record's type and identifier, and the field."""

    entity_type: str
    identifier: str
    field: str


def holders(facade: Any, child_type: str) -> list[Holder]:
    """Every field, of any entity type, that holds ``child_type``, in profile order."""
    found: list[Holder] = []
    for entity_type in getattr(facade, "entities", []) or []:
        helper = getattr(facade, entity_type, None)
        for field, target in (getattr(helper, "child_fields", {}) or {}).items():
            if target == child_type:
                found.append(Holder(entity_type, field))
    return found


def _label(holder: Holder, all_holders: list[Holder]) -> str:
    """``Type``, or ``Type.field`` when that type holds the row in two fields."""
    same_type = [h for h in all_holders if h.entity_type == holder.entity_type]
    return (
        f"{holder.entity_type}.{holder.field}"
        if len(same_type) > 1
        else holder.entity_type
    )


def value_for(facade: Any, child_type: str, placement: Placement) -> str:
    """The cell text naming ``placement`` for a row of ``child_type``."""
    label = _label(
        Holder(placement.entity_type, placement.field), holders(facade, child_type)
    )
    return f"{label}: {placement.identifier}"


def parse(facade: Any, child_type: str, value: str) -> Placement:
    """Read a **Belongs to** cell.

    Args:
        facade: The dataset's profile facade.
        child_type: The entity type of the row's sheet.
        value: The cell text.

    Returns:
        The holding record's type and identifier, and the holding field.

    Raises:
        ValueError: With what may be written, when the value is not
            ``Type: identifier`` or names a type that does not hold the row,
            or holds it in more than one field without naming which.
    """
    label, sep, identifier = value.partition(": ")
    identifier = identifier.strip()
    candidates = holders(facade, child_type)
    allowed = sorted({_label(h, candidates) for h in candidates})
    if not sep or not identifier:
        raise ValueError(
            f"'{value}' is not 'Type: identifier'; write one of {', '.join(allowed)} "
            "followed by ': ' and the record's identifier"
        )
    entity_type, _, field = label.strip().partition(".")
    matching = [
        h
        for h in candidates
        if h.entity_type == entity_type and (not field or h.field == field)
    ]
    if not matching:
        raise ValueError(
            f"{label.strip()} does not hold {child_type}; write one of {', '.join(allowed)}"
        )
    if len(matching) > 1:
        raise ValueError(
            f"{entity_type} holds {child_type} in {', '.join(h.field for h in matching)}; "
            f"write which, as {entity_type}.<field>: {identifier}"
        )
    return Placement(entity_type, identifier, matching[0].field)


def apply_belongs_to_validation(
    wb: Workbook,
    facade: Any,
    sheets: dict[str, Worksheet],
    columns_by_entity: dict[str, list[str]],
) -> None:
    """Give every **Belongs to** column a dropdown of the records that may hold the row.

    The entries are formulas over the holders' identifier columns, written into
    a hidden sheet, so a parent row added in the same edit appears in its
    children's dropdown. A blank holder row yields a blank entry.
    """
    parents: Worksheet | None = None
    column = 0
    for child_type, columns in columns_by_entity.items():
        ws = sheets.get(child_type)
        if ws is None or BELONGS_TO not in columns:
            continue
        child_holders = holders(facade, child_type)
        sources = []
        for holder in child_holders:
            helper = getattr(facade, holder.entity_type, None)
            identifier = getattr(helper, "identifier_field", None)
            holder_columns = columns_by_entity.get(holder.entity_type) or []
            if holder.entity_type in sheets and identifier in holder_columns:
                letter = get_column_letter(holder_columns.index(identifier) + 1)
                sources.append(
                    (
                        _label(holder, child_holders),
                        sheets[holder.entity_type].title,
                        letter,
                    )
                )
        if not sources:
            continue
        if parents is None:
            parents = wb.create_sheet(PARENTS_SHEET)
            parents.sheet_state = "hidden"
        column += 1
        out = get_column_letter(column)
        parents.cell(row=1, column=column, value=child_type)
        row = 2
        for label, sheet, letter in sources:
            for source_row in range(2, PARENT_ROWS + 2):
                cell = f"'{sheet}'!${letter}${source_row}"
                parents.cell(
                    row=row,
                    column=column,
                    value=f'=IF({cell}="","","{label}: "&{cell})',
                )
                row += 1
        target = get_column_letter(columns.index(BELONGS_TO) + 1)
        validation = DataValidation(
            type="list",
            formula1=f"'{PARENTS_SHEET}'!${out}$2:${out}${row - 1}",
            allow_blank=True,
            showErrorMessage=True,
            showInputMessage=True,
            errorStyle="warning",
            errorTitle="Belongs to: no such record",
            error=(
                "Choose the record this row belongs to from the list. Continue "
                "anyway if that record is still to be added."
            ),
            promptTitle=BELONGS_TO,
            prompt=f"The record that holds this {child_type}.",
        )
        ws.add_data_validation(validation)
        validation.add(f"{target}2:{target}{PARENT_ROWS + 1}")
