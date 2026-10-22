"""The workbook's first sheet: how to add data to it.

A person given an exported dataset had to discover that a table grows when one
types below it, which column places a row in the tree, and where a field's
rules are written. This sheet says so before any data sheet is opened.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.styles import Alignment, Font

from metaseed.ui.services.belongs_to import BELONGS_TO

if TYPE_CHECKING:
    from openpyxl import Workbook

INSTRUCTIONS_SHEET = "How to add data"

STEPS = (
    (
        "Each sheet is one kind of record",
        "Every sheet after this one holds the records of one entity type, one per row. "
        "The first row names the fields.",
    ),
    (
        "Add a row",
        "Type in the first empty row directly below a sheet's table. The table grows to "
        "include the row, and its dropdowns and checks apply to it.",
    ),
    (
        f"Say where it belongs: {BELONGS_TO}",
        f"The {BELONGS_TO} column places the row in the dataset. Choose the record that "
        "holds it from the dropdown, for example 'Study: STU-1'. A record added in the "
        "same edit appears in the list once its identifier is filled in. Leave the cell "
        "empty only on the sheet of the top-level record.",
    ),
    (
        "Fill in the fields",
        "Hover over a column heading to read what the field holds: whether it is "
        "required, its type, an example, and its pattern, allowed values and rules. A "
        "column with allowed values has a dropdown. A cell that breaks a rule is coloured.",
    ),
    (
        "Do not change",
        "Leave the sheets' names and first rows as they are, and the hidden sheets "
        "alone: the import reads them to know where each value goes.",
    ),
    (
        "Import",
        "Import the saved workbook into the dataset. It is read whole: a row whose "
        f"{BELONGS_TO} names no record is reported with its sheet and row, and nothing "
        "is changed until every row can be placed.",
    ),
)


def write_instructions(wb: Workbook) -> None:
    """Write the instructions as the workbook's first sheet."""
    ws = wb.create_sheet(INSTRUCTIONS_SHEET, 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    ws.append(["How to add data to this workbook"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for step, text in STEPS:
        ws.append([step, text])
        row = ws.max_row
        ws.cell(row=row, column=1).font = Font(bold=True)
        for column in (1, 2):
            ws.cell(row=row, column=column).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
